import base64, io, json, os, re, uuid
from pathlib import Path

import numpy as np
from dotenv import load_dotenv
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.responses import FileResponse
from pydantic import BaseModel

load_dotenv()
MODEL = os.getenv("OPENAI_MODEL", "gpt-5.5")
DATA = Path(os.getenv("DATA_DIR", "data"))
DATA.mkdir(parents=True, exist_ok=True)
INDEX, VEC = DATA / "index.json", DATA / "vectors.npy"
DIM = 384  # bge-small-en-v1.5

app = FastAPI(title="Cloud RAG")
_client = _embedder = None


def openai_client():
    global _client
    if _client is None:
        from openai import OpenAI
        _client = OpenAI()  # reads OPENAI_API_KEY
    return _client


def embed(texts):
    global _embedder
    if _embedder is None:
        from fastembed import TextEmbedding
        _embedder = TextEmbedding("BAAI/bge-small-en-v1.5")
    v = np.array(list(_embedder.embed(texts)), dtype="float32")
    return v / np.linalg.norm(v, axis=1, keepdims=True)


def load():
    if INDEX.exists() and VEC.exists():
        return json.loads(INDEX.read_text()), np.load(VEC)
    return [], np.zeros((0, DIM), dtype="float32")


def save(chunks, vecs):
    INDEX.write_text(json.dumps(chunks))
    np.save(VEC, vecs)


def describe_image(ext, data):
    mt = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
          ".gif": "image/gif", ".webp": "image/webp"}[ext]
    image_data = f"data:{mt};base64,{base64.b64encode(data).decode()}"
    r = openai_client().responses.create(
        model=MODEL,
        input=[{
            "role": "user",
            "content": [
                {"type": "input_text", "text": "Transcribe all text in this image, then describe it in detail."},
                {"type": "input_image", "image_url": image_data, "detail": "auto"},
            ],
        }]
    )
    return r.output_text


def extract(name, data):
    ext = Path(name).suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        return "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(data)).pages)
    if ext == ".docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        rows = [" | ".join(c.text for c in r.cells) for t in d.tables for r in t.rows]
        return "\n".join([p.text for p in d.paragraphs] + rows)
    if ext == ".pptx":
        from pptx import Presentation
        out = []
        for i, s in enumerate(Presentation(io.BytesIO(data)).slides, 1):
            out.append(f"[Slide {i}]")
            out += [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
        return "\n".join(out)
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        out = []
        for ws in load_workbook(io.BytesIO(data), read_only=True, data_only=True):
            out.append(f"[Sheet {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                if any(c is not None for c in row):
                    out.append(" | ".join("" if c is None else str(c) for c in row))
        return "\n".join(out)
    if ext in (".html", ".htm"):
        from bs4 import BeautifulSoup
        return BeautifulSoup(data, "html.parser").get_text("\n")
    if ext in (".png", ".jpg", ".jpeg", ".gif", ".webp"):
        return describe_image(ext, data)
    text = data.decode("utf-8", errors="ignore")  # txt, md, csv, json, code, logs...
    if "\x00" in text[:2000]:
        raise ValueError("unsupported binary file")
    return text


def chunk(text, size=1000, overlap=150):
    text = re.sub(r"[ \t]+", " ", re.sub(r"\n{3,}", "\n\n", text)).strip()
    return [text[i:i + size] for i in range(0, len(text), size - overlap) if text[i:i + size].strip()]


@app.post("/api/upload")
def upload(files: list[UploadFile] = File(...)):
    chunks, vecs, added = *load(), []
    for f in files:
        try:
            pieces = chunk(extract(f.filename, f.file.read()))
        except Exception as e:
            raise HTTPException(422, f"{f.filename}: {e}")
        if not pieces:
            raise HTTPException(422, f"{f.filename}: no readable text found")
        doc_id = uuid.uuid4().hex[:8]
        chunks += [{"doc_id": doc_id, "name": f.filename, "n": i + 1, "text": p}
                   for i, p in enumerate(pieces)]
        vecs = np.vstack([vecs, embed(pieces)])
        added.append({"id": doc_id, "name": f.filename, "chunks": len(pieces)})
    save(chunks, vecs)
    return added


@app.get("/api/documents")
def documents():
    docs = {}
    for c in load()[0]:
        d = docs.setdefault(c["doc_id"], {"id": c["doc_id"], "name": c["name"], "chunks": 0})
        d["chunks"] += 1
    return list(docs.values())


@app.delete("/api/documents/{doc_id}")
def delete(doc_id: str):
    chunks, vecs = load()
    keep = np.array([i for i, c in enumerate(chunks) if c["doc_id"] != doc_id], dtype=int)
    save([chunks[i] for i in keep], vecs[keep])
    return {"ok": True}


class Ask(BaseModel):
    question: str
    use_web: bool = True
    history: list[dict] = []


SYSTEM = """You answer questions using the user's uploaded file excerpts and, when available, web search.
- Prefer the file excerpts. Cite them inline like [F1], [F2].
- Use web search only if the files don't cover the question or the user asks for external or current information. Start those parts with "From the web:".
- If neither source has the answer, say so plainly. Never invent facts."""


@app.post("/api/ask")
def ask(q: Ask):
    chunks, vecs = load()
    ctx, sources = [], []
    if chunks:
        scores = vecs @ embed([q.question])[0]
        for rank, i in enumerate(np.argsort(-scores)[:5], 1):
            c = chunks[i]
            ctx.append(f"[F{rank}] {c['name']}, part {c['n']}\n{c['text']}")
            sources.append({"title": f"F{rank} · {c['name']} (part {c['n']})"})
    excerpts = "\n\n".join(ctx) or "No files uploaded."

    input_items = []
    for h in q.history[-6:]:
        role = h.get("role")
        content = h.get("content", "")
        if role in ("user", "assistant"):
            input_items.append({"role": role, "content": content})
    input_items.append({
        "role": "user",
        "content": f"<file_excerpts>\n{excerpts}\n</file_excerpts>\n\nQuestion: {q.question}"
    })

    kwargs = {
        "model": MODEL,
        "instructions": SYSTEM,
        "input": input_items,
    }
    if q.use_web:
        kwargs["tools"] = [{"type": "web_search", "search_context_size": "low"}]

    try:
        r = openai_client().responses.create(**kwargs)
    except Exception as e:
        raise HTTPException(502, str(e))

    answer = r.output_text
    # Keep the UI's source list useful. OpenAI's Responses API exposes web-search
    # citations as annotations on output text; extract their titles/URLs when present.
    for item in getattr(r, "output", []):
        if getattr(item, "type", None) == "message":
            for content in getattr(item, "content", []):
                for ann in getattr(content, "annotations", []) or []:
                    url = getattr(ann, "url", None)
                    title = getattr(ann, "title", None) or url
                    if url:
                        sources.append({"title": title, "url": url})
    return {"answer": answer, "sources": sources}


@app.get("/")
def home():
    return FileResponse(Path(__file__).parent / "static" / "index.html")
